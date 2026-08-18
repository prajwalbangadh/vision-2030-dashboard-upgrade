from __future__ import annotations
import argparse, json, os, re, threading, webbrowser
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from openpyxl import load_workbook

SITE=Path(__file__).resolve().parent
PROJECT=SITE.parent
OUTPUT=PROJECT/'output'
TEMPLATES=SITE/'templates'
PAGES=('briefing.html','insights.html','data_explorer.html')
ALIASES={
 'Leader Effectiveness':'Leader Effectiveness (Facility Leadership)',
 'Opportunity to Learn and Grow':'Opportunity to Learn and Grow',
 'Forecast and Hire':'Forecast and Hire',
 'Team Member Engagement':'Team Member Engagement',
 'Total Turnover (Rolling 12)':'Total Turnover',
 'Physician Engagement':'Physician Engagement',
 'Annual Active Users':'Digital Tool Utilization',
 'LTR ED':'Patient Experience – ED LTR',
 'LTR Inpatient':'Patient Experience – Inpatient LTR',
 'LTR Med Practice':'Patient Experience – Med Practice LTR',
 'CMS Star Rating':'CMS Overall Hospital Star Rating',
 'Leapfrog Safety Grade':'Leapfrog Hospital Safety Grade',
 'All-Adult Inpatient Mortality':'All-Adult Inpatient Mortality',
 'Length of Stay O/E':'Length of Stay O/E',
 'EBITDA Dollars':'EBITDA $',
 'TOR Growth Rate':'TOR Growth Rate',
 'Domestic Spend':'Domestic Spend',
 'Internal Leader Fill Rate':'Internal Leader Fill Rate',
}
DIV_SCOPE={'CFD':'CFD','EFD':'EFD','WFD':'WFD','PHD':'PHD','MSD':'MSD','MSD 1':'MSD','MSD 2':'MSD'}

def latest(root:Path)->Path:
 files=[p for p in root.glob('*/vision2030_business_table_*.xlsx') if p.is_file() and not p.name.startswith('~$')]
 if not files: raise FileNotFoundError(f'No Vision 2030 business workbook found under {root}')
 stamped=[p for p in files if re.fullmatch(r'\d{4}-\d{2}-\d{2}_\d{6}',p.parent.name)]
 return max(stamped or files,key=lambda p:(p.parent.name,p.name))

def rows(ws):
 it=ws.iter_rows(values_only=True); hdr=[str(v).strip() if v is not None else '' for v in next(it)]
 for vals in it:
  if any(v is not None and str(v).strip() for v in vals): yield {hdr[i]:vals[i] if i<len(vals) else None for i in range(len(hdr))}

def text(v):
 return 'N/A' if v is None or not str(v).strip() else str(v).strip()

def score(v):
 v=text(v).lower(); return v if v in {'green','red','gray'} else 'gray'

def period(summary):
 labels=[text(r.get('Data As Of')) for r in summary if text(r.get('Data As Of'))!='N/A']
 return max(labels,key=labels.count) if labels else 'Latest extraction'

def extract_model(html):
 marker='const VD ='; start=html.index(marker)+len(marker)
 model,used=json.JSONDecoder().raw_decode(html[start:])
 return model,start,start+used

def update_scope(model,scope,metric,value,status):
 block=model.get('metricData',{}).get(scope)
 if not block: return False
 block.setdefault('values',{})[metric]=value
 block.setdefault('scores',{})[metric]=status
 return True

def build_model(template_model,book):
 model=template_model
 summary=list(rows(book['Vision 2030']))
 model['period']=period(summary)
 updated=0
 for r in summary:
  metric=ALIASES.get(text(r.get('Metric')))
  if metric and update_scope(model,'SYS',metric,text(r.get('Value')),score(r.get('Variance Color'))): updated+=1
 # Direct division/location rows only. Region/facility structure remains intact until explicitly mapped.
 for sheet in ('Learning','Team','Consumer','Clinical','Financial','Risk','WPC'):
  for r in rows(book[sheet]):
   metric=ALIASES.get(text(r.get('Metric')))
   scope=DIV_SCOPE.get(text(r.get('Division')))
   if metric and scope and text(r.get('Location Type')).lower()=='division':
    if update_scope(model,scope,metric,text(r.get('Value')),score(r.get('Variance Color'))): updated+=1
 model['_source']={'workbook':str(WORKBOOK),'updated_fields':updated}
 return model,updated

def inject(template_path,out_path,book):
 html=template_path.read_text(encoding='utf-8')
 model,start,end=extract_model(html)
 model,updated=build_model(model,book)
 payload=json.dumps(model,ensure_ascii=False,separators=(',',':'))
 rendered=html[:start]+payload+html[end:]
 temp=out_path.with_suffix('.tmp'); temp.write_text(rendered,encoding='utf-8'); temp.replace(out_path)
 # Verify injected object is valid and UI markup outside VD stays byte-for-byte identical.
 check,cs,ce=extract_model(rendered)
 assert check['_source']['workbook']==str(WORKBOOK)
 assert html[:start]==rendered[:cs] and html[end:]==rendered[ce:]
 return updated

class Handler(SimpleHTTPRequestHandler):
 def end_headers(self): self.send_header('Cache-Control','no-store'); super().end_headers()

def serve(port):
 os.chdir(SITE); server=None
 for p in range(port,port+11):
  try: server=ThreadingHTTPServer(('127.0.0.1',p),Handler); port=p; break
  except OSError: pass
 if server is None: raise OSError('No available local port')
 url=f'http://127.0.0.1:{port}/briefing.html'; print('Website:',url); print('Press Ctrl+C to stop.')
 threading.Timer(.7,lambda:webbrowser.open(url)).start()
 try: server.serve_forever()
 except KeyboardInterrupt: print('\nStopped.')
 finally: server.server_close()

def main():
 global WORKBOOK
 ap=argparse.ArgumentParser(); ap.add_argument('--workbook',type=Path); ap.add_argument('--output-root',type=Path,default=OUTPUT); ap.add_argument('--port',type=int,default=8000); ap.add_argument('--build-only',action='store_true'); a=ap.parse_args()
 WORKBOOK=(a.workbook.resolve() if a.workbook else latest(a.output_root.resolve()))
 print('Latest workbook selected:',WORKBOOK)
 book=load_workbook(WORKBOOK,read_only=True,data_only=True)
 required={'Vision 2030','Learning','Team','Consumer','Clinical','Financial','Risk','WPC'}
 missing=required-set(book.sheetnames)
 if missing: raise ValueError('Missing sheets: '+', '.join(sorted(missing)))
 total=0
 for page in PAGES:
  count=inject(TEMPLATES/page,SITE/page,book); total+=count; print(f'Built {page}: {count} factual fields updated')
 print('Contemporary UI preserved; total page-field updates:',total)
 if not a.build_only: serve(a.port)
if __name__=='__main__': main()
