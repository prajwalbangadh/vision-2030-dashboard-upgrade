import csv
import importlib
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_FOLDER = Path(__file__).resolve().parent

DATALINK_PAGE_URL = (
    "https://datalink.adventhealth.com/Home/vision2030"
)

API_URL = (
    "https://datalink.adventhealth.com/api/metrics/vision2030"
)

TIMESTAMP = datetime.now().strftime("%Y-%m-%d_%H%M%S")

OUTPUT_FOLDER = PROJECT_FOLDER / "output" / TIMESTAMP

JSON_OUTPUT_FILE = (
    OUTPUT_FOLDER
    / f"vision2030_raw_{TIMESTAMP}.json"
)

RAW_CSV_OUTPUT_FILE = (
    OUTPUT_FOLDER
    / f"vision2030_metrics_raw_{TIMESTAMP}.csv"
)

BUSINESS_EXCEL_OUTPUT_FILE = (
    OUTPUT_FOLDER
    / f"vision2030_business_table_{TIMESTAMP}.xlsx"
)

# ============================================================
# EXACT VISION 2030 EXECUTIVE TABLE CONFIGURATION
# ============================================================
#
# The API provides changing values.
#
# This configuration explicitly controls:
# - the 20 displayed rows
# - category
# - aspiration
# - API metric selector
# - value format
# - goal display
# - variance source and display
# - YTD handling
# - output order
#
# Goal operators are NOT inferred from the API.
# Goal text is explicitly hardcoded per metric.
# ============================================================

METRIC_RULES = [
    {
        "order": 1,
        "category": "LEARNING",
        "aspiration": "Dynamic Learning Community",
        "display_name": "Internal Leader Fill Rate",
        "metric_id": 84,
        "api_names": [
            "Internal Leader Fill Rate",
        ],
        "value_format": "percent_1",
        "goal_display": "N/A",
        "variance_method": "none",
        "variance_format": "none",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 2,
        "category": "LEARNING",
        "aspiration": "Dynamic Learning Community",
        "display_name": "Leader Effectiveness",
        "metric_id": 77,
        "api_names": [
            "Leader Effectiveness",
            "Leader Effectiveness Improvement over prior score",
        ],
        "value_format": "decimal_1",
        "goal_display": "≥ 81",
        "variance_method": "api_rounded",
        "variance_format": "percent_1",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 3,
        "category": "LEARNING",
        "aspiration": "Dynamic Learning Community",
        "display_name": "Opportunity to Learn and Grow",
        "metric_id": 75,
        "api_names": [
            "Opportunity to Learn and Grow",
            "Opp to Learn and Grow (PeopleSoft)",
        ],
        "value_format": "integer",
        "goal_display": "≥ 79",
        "variance_method": "api_raw",
        "variance_format": "percent_2",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 4,
        "category": "TEAM",
        "aspiration": "Team Member Promise",
        "display_name": "Forecast and Hire",
        "metric_id": 87,
        "api_names": [
            "Forecast and Hire",
        ],
        "value_format": "integer_comma",
        "goal_display": "≥ 12,176",
        "variance_method": "api_rounded",
        "variance_format": "percent_0",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 5,
        "category": "TEAM",
        "aspiration": "Team Member Promise",
        "display_name": "Team Member Engagement",
        "metric_id": 76,
        "api_names": [
            "Team Member Engagement",
            "Glint",
        ],
        "value_format": "integer",
        "goal_display": "≥ 81",
        "variance_method": "api_raw",
        "variance_format": "percent_2",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 6,
        "category": "TEAM",
        "aspiration": "Team Member Promise",
        "display_name": "Total Turnover (Rolling 12)",
        "metric_id": 71,
        "api_names": [
            "Total Turnover (Rolling 12)",
        ],
        "value_format": "percent_1",
        "goal_display": "≤ 17.3%",
        "variance_method": "hardcoded_absolute",
        "variance_goal": 17.3,
        "variance_sign": "goal_minus_value",
        "variance_format": "percent_1",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 7,
        "category": "TEAM",
        "aspiration": "Team Member Promise",
        "display_name": "Physician Engagement",
        "metric_id": None,
        "api_names": [
            "Physician Engagement",
        ],
        "value_format": "integer",
        "goal_display": "≥ 75",
        "variance_method": "hardcoded_relative",
        "variance_goal": 75,
        "variance_sign": "value_minus_goal",
        "variance_format": "percent_0",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 8,
        "category": "CONSUMER",
        "aspiration": "Consumer Focused Connected Network",
        "display_name": "Digital Tool Utilization",
        "metric_id": None,
        "api_names": [
            "Digital Tool Utilization",
        ],
        "value_format": "millions_1",
        "goal_display": "≥ 3.6M",
        "variance_method": "hardcoded_relative",
        "variance_goal": 3_600_000,
        "variance_sign": "value_minus_goal",
        "variance_format": "percent_1",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 9,
        "category": "CONSUMER",
        "aspiration": "Consumer Focused Connected Network",
        "display_name": "LTR ED",
        "metric_id": 20,
        "api_names": [
            "LTR ED",
            "CNS ED LTR",
        ],
        "value_format": "integer",
        "goal_display": "≥ 75",
        "variance_method": "hardcoded_relative",
        "variance_goal": 75,
        "variance_sign": "value_minus_goal",
        "variance_format": "percent_0",
        "ytd_format": "integer",
        "selection": "corporate",
    },
    {
        "order": 10,
        "category": "CONSUMER",
        "aspiration": "Consumer Focused Connected Network",
        "display_name": "LTR Inpatient",
        "metric_id": 39,
        "api_names": [
            "LTR Inpatient",
            "CNS LTR Inpatient",
        ],
        "value_format": "integer",
        "goal_display": "≥ 75",
        "variance_method": "hardcoded_relative",
        "variance_goal": 75,
        "variance_sign": "value_minus_goal",
        "variance_format": "percent_1",
        "ytd_format": "integer",
        "selection": "corporate",
    },
    {
        "order": 11,
        "category": "CONSUMER",
        "aspiration": "Consumer Focused Connected Network",
        "display_name": "LTR Med Practice",
        "metric_id": 40,
        "api_names": [
            "LTR Med Practice",
            "CNS LTR Med Practice",
            "LTR Medical Practice",
        ],
        "value_format": "integer",
        "goal_display": "≥ 75",
        "variance_method": "hardcoded_relative",
        "variance_goal": 75,
        "variance_sign": "value_minus_goal",
        "variance_format": "percent_1",
        "ytd_format": "integer",
        "selection": "corporate",
    },
    {
        "order": 12,
        "category": "CLINICAL",
        "aspiration": "Clinical Excellence",
        "display_name": "Length of Stay O/E",
        "metric_id": None,
        "api_names": [
            "Length of Stay O/E",
            "Adult Length of Stay O/E",
        ],
        "value_format": "coming_soon",
        "goal_display": "N/A",
        "variance_method": "none",
        "variance_format": "none",
        "ytd_format": "none",
        "selection": "placeholder",
    },
    {
        "order": 13,
        "category": "CLINICAL",
        "aspiration": "Clinical Excellence",
        "display_name": "CMS Star Rating",
        "metric_id": 10,
        "api_names": [
            "CMS Star Rating",
        ],
        "value_format": "percent_1",
        "goal_display": "4 & 5 Stars",
        "variance_method": "none",
        "variance_format": "none",
        "ytd_format": "none",
        "selection": "cms_star_summary",
    },
    {
        "order": 14,
        "category": "CLINICAL",
        "aspiration": "Clinical Excellence",
        "display_name": "Leapfrog Safety Grade",
        "metric_id": 42,
        "api_names": [
            "Leapfrog Safety Grade",
        ],
        "value_format": "percent_1",
        "goal_display": "A's",
        "variance_method": "none",
        "variance_format": "none",
        "ytd_format": "none",
        "selection": "leapfrog_summary",
    },
    {
        "order": 15,
        "category": "FINANCIAL",
        "aspiration": "Financial Strength and Growth",
        "display_name": "EBITDA Dollars",
        "metric_id": None,
        "api_names": [
            "EBITDA Dollars",
        ],
        "value_format": "currency_millions_1",
        "goal_display": "≥ $279.5M",
        "variance_method": "hardcoded_relative",
        "variance_goal": 279_500_000,
        "variance_sign": "value_minus_goal",
        "variance_format": "percent_1",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 16,
        "category": "FINANCIAL",
        "aspiration": "Financial Strength and Growth",
        "display_name": "TOR Dollars",
        "metric_id": None,
        "api_names": [
            "TOR Dollars",
        ],
        "value_format": "currency_billions_1",
        "goal_display": "≥ $2B",
        "variance_method": "api_rounded",
        "variance_format": "percent_1",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 17,
        "category": "FINANCIAL",
        "aspiration": "Financial Strength and Growth",
        "display_name": "TOR Growth Rate",
        "metric_id": 68,
        "api_names": [
            "TOR Growth Rate",
        ],
        "value_format": "percent_1",
        "goal_display": "≥ 12.3%",
        "variance_method": "hardcoded_absolute",
        "variance_goal": 12.3,
        "variance_sign": "value_minus_goal",
        "variance_format": "percent_1",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 18,
        "category": "RISK",
        "aspiration": "Managed Population Risk",
        "display_name": "Domestic Spend",
        "metric_id": None,
        "api_names": [
            "Domestic Spend",
        ],
        "value_format": "percent_1",
        "goal_display": "≥ 76.6%",
        "variance_method": "hardcoded_relative",
        "variance_goal": 76.6,
        "variance_sign": "value_minus_goal",
        "variance_format": "percent_1",
        "ytd_format": "none",
        "selection": "corporate",
    },
    {
        "order": 19,
        "category": "WPC",
        "aspiration": "Whole-Person Care",
        "display_name": "Community Impact",
        "metric_id": None,
        "api_names": [
            "Community Impact",
        ],
        "value_format": "coming_soon",
        "goal_display": "N/A",
        "variance_method": "none",
        "variance_format": "none",
        "ytd_format": "none",
        "selection": "placeholder",
    },
    {
        "order": 20,
        "category": "WPC",
        "aspiration": "Whole-Person Care",
        "display_name": "Mission Integration Plan",
        "metric_id": None,
        "api_names": [
            "Mission Integration Plan",
        ],
        "value_format": "coming_soon",
        "goal_display": "N/A",
        "variance_method": "none",
        "variance_format": "none",
        "ytd_format": "none",
        "selection": "placeholder",
    },
]

CATEGORY_SHEETS = {
    "LEARNING": {
        "sheet_name": "Learning",
        "aspiration": "Dynamic Learning Community",
    },
    "TEAM": {
        "sheet_name": "Team",
        "aspiration": "Team Member Promise",
    },
    "CONSUMER": {
        "sheet_name": "Consumer",
        "aspiration": "Consumer Focused Connected Network",
    },
    "CLINICAL": {
        "sheet_name": "Clinical",
        "aspiration": "Clinical Excellence",
    },
    "FINANCIAL": {
        "sheet_name": "Financial",
        "aspiration": "Financial Strength and Growth",
    },
    "RISK": {
        "sheet_name": "Risk",
        "aspiration": "Managed Population Risk",
    },
    "WPC": {
        "sheet_name": "WPC",
        "aspiration": "Whole-Person Care",
    },
}

DISPLAY_LEVELS = {
    0: "Executive",
    1: "First Drilldown",
    2: "Final Drilldown",
}

CATEGORY_SHEET_COLUMNS = [
    "Metric",
    "Level",
    "Division",
    "Location Type",
    "Location",
    "Value",
    "Goal",
    "Variance Direction",
    "Variance",
    "Variance Color",
    "YTD",
    "Data As Of",
]

METRIC_DISPLAY_LEVEL_RULES = {
    "Internal Leader Fill Rate": {
        "Corporate": 0,
        "Division": 1,
        "Region": 2,
    },
    "Domestic Spend": {
        "Corporate": 0,
        "Division": 1,
        "Region": 1,
    },
    "EBITDA Dollars": {
        "Corporate": 0,
        "Division": 1,
        "Facility": 2,
    },
    "CMS Star Rating": {
        "Corporate": 0,
        "Facility": 1,
    },
    "Leapfrog Safety Grade": {
        "Corporate": 0,
        "Campus": 1,
    },
}

# ============================================================
# PACKAGE INSTALLATION
# ============================================================

def ensure_pip():
    try:
        subprocess.run(
            [
                sys.executable,
                "-m",
                "pip",
                "--version",
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )

    except (subprocess.CalledProcessError, FileNotFoundError):
        print("pip is not available. Setting up pip...")

        subprocess.check_call(
            [
                sys.executable,
                "-m",
                "ensurepip",
                "--upgrade",
            ]
        )


def ensure_package(package_name):
    try:
        importlib.import_module(package_name)
        print(f"{package_name} is already installed.")

    except ModuleNotFoundError:
        print(
            f"{package_name} is not installed. "
            "Installing it now..."
        )

        ensure_pip()

        subprocess.check_call(
            [
                sys.executable,
                "-m",
                "pip",
                "install",
                "--upgrade",
                package_name,
            ]
        )

        importlib.invalidate_caches()
        importlib.import_module(package_name)

        print(f"{package_name} installed successfully.")

ensure_package("selenium")
ensure_package("openpyxl")


from selenium import webdriver
from selenium.webdriver.edge.options import Options
from selenium.webdriver.support.ui import WebDriverWait

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# GENERAL HELPERS
# ============================================================

def safe_float(value):
    if value in (None, "", "N/A"):
        return None

    try:
        return float(value)

    except (TypeError, ValueError):
        return None


def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


def get_record_names(record):
    names = {
        clean_text(record.get("metric_name")),
        clean_text(record.get("mtrc_cmn_nm")),
    }

    return {
        name.casefold()
        for name in names
        if name
    }


def rule_matches_record(rule, record):
    record_metric_id = record.get("mtrc_id")
    configured_metric_id = rule.get("metric_id")

    if (
        configured_metric_id is not None
        and record_metric_id == configured_metric_id
    ):
        return True

    record_names = get_record_names(record)

    configured_names = {
        clean_text(name).casefold()
        for name in rule.get("api_names", [])
        if clean_text(name)
    }

    return bool(record_names.intersection(configured_names))


def is_adventhealth_corporate_record(record):
    return (
        clean_text(record.get("rollup_type")).casefold()
        == "corporate"
        and clean_text(record.get("location_name")).casefold()
        == "adventhealth"
    )


def get_primary_value(record):
    rounded_value = record.get(
        "prim_org_mtrc_calc_rnd_nbr"
    )

    if rounded_value is not None:
        return safe_float(rounded_value)

    return safe_float(record.get("metric_value"))


def get_ytd_value(record):
    rounded_ytd_value = record.get(
        "prim_org_rcnt_year_mtrc_calc_rnd_nbr"
    )

    if rounded_ytd_value is not None:
        return safe_float(rounded_ytd_value)

    return safe_float(
        record.get("prim_org_rcnt_year_calc_lbl")
    )


# ============================================================
# BROWSER EXTRACTION
# ============================================================

def get_json_text_from_browser(driver):
    script = """
    const preElement = document.querySelector("pre");

    if (preElement && preElement.textContent.trim()) {
        return preElement.textContent.trim();
    }

    if (document.body && document.body.innerText.trim()) {
        return document.body.innerText.trim();
    }

    return "";
    """

    return driver.execute_script(script)


def validate_records(records):
    if not isinstance(records, list):
        raise TypeError(
            "Expected the API response to be a JSON list, "
            f"but received {type(records).__name__}."
        )

    if not records:
        raise ValueError(
            "The Vision 2030 API returned zero records."
        )

    required_fields = {
        "application",
        "grouper",
        "metric_name",
        "metric_value",
        "rollup_type",
        "location_name",
    }

    first_record = records[0]

    if not isinstance(first_record, dict):
        raise TypeError(
            "Expected each API record to be a JSON object."
        )

    missing_fields = required_fields.difference(
        first_record.keys()
    )

    if missing_fields:
        raise ValueError(
            "The first API record is missing required fields: "
            + ", ".join(sorted(missing_fields))
        )


# ============================================================
# RAW OUTPUT
# ============================================================

def get_all_columns(records):
    columns = []
    seen_columns = set()

    for record in records:
        if not isinstance(record, dict):
            continue

        for column_name in record.keys():
            if column_name not in seen_columns:
                seen_columns.add(column_name)
                columns.append(column_name)

    return columns


def save_json(records):
    JSON_OUTPUT_FILE.write_text(
        json.dumps(
            records,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def save_raw_csv(records):
    columns = get_all_columns(records)

    with RAW_CSV_OUTPUT_FILE.open(
        mode="w",
        newline="",
        encoding="utf-8-sig",
    ) as csv_file:

        writer = csv.DictWriter(
            csv_file,
            fieldnames=columns,
            extrasaction="ignore",
        )

        writer.writeheader()
        writer.writerows(records)

    return columns


# ============================================================
# VALUE FORMATTING
# ============================================================

def format_value(value, format_type):
    if format_type == "coming_soon":
        return "Coming Soon"

    numeric_value = safe_float(value)

    if numeric_value is None:
        return "N/A"

    if format_type == "integer":
        return f"{numeric_value:,.0f}"

    if format_type == "integer_comma":
        return f"{numeric_value:,.0f}"

    if format_type == "decimal_1":
        return f"{numeric_value:,.1f}"

    if format_type == "decimal_2":
        return f"{numeric_value:,.2f}"

    if format_type == "percent_1":
        return f"{numeric_value:,.1f}%"

    if format_type == "millions_1":
        return f"{numeric_value / 1_000_000:,.1f}M"

    if format_type == "currency_millions_1":
        return (
            f"${numeric_value / 1_000_000:,.1f}M"
        )

    if format_type == "currency_billions_1":
        return (
            f"${numeric_value / 1_000_000_000:,.1f}B"
        )

    return str(value)


# ============================================================
# RECORD SELECTION
# ============================================================

def find_corporate_record(records, rule):
    matching_records = [
        record
        for record in records
        if rule_matches_record(rule, record)
    ]

    corporate_matches = [
        record
        for record in matching_records
        if is_adventhealth_corporate_record(record)
    ]

    if corporate_matches:
        return corporate_matches[0]

    adventhealth_matches = [
        record
        for record in matching_records
        if (
            clean_text(
                record.get("location_name")
            ).casefold()
            == "adventhealth"
        )
    ]

    if adventhealth_matches:
        return adventhealth_matches[0]

    return None


def get_matching_records(records, rule):
    return [
        record
        for record in records
        if rule_matches_record(rule, record)
    ]


# ============================================================
# SPECIAL CLINICAL SUMMARY METRICS
# ============================================================

def calculate_cms_star_summary(records, rule):
    matching_records = get_matching_records(
        records,
        rule,
    )

    facility_records = [
        record
        for record in matching_records
        if clean_text(
            record.get("rollup_type")
        ).casefold()
        == "facility"
        and get_primary_value(record) is not None
    ]

    if not facility_records:
        return None

    four_or_five_star_count = sum(
        1
        for record in facility_records
        if get_primary_value(record) in (4, 5)
    )

    return (
        four_or_five_star_count
        / len(facility_records)
        * 100
    )


def calculate_leapfrog_summary(records, rule):
    matching_records = get_matching_records(
        records,
        rule,
    )

    campus_records = [
        record
        for record in matching_records
        if clean_text(
            record.get("rollup_type")
        ).casefold()
        == "campus"
        and get_primary_value(record) is not None
    ]

    if not campus_records:
        return None

    grade_a_count = sum(
        1
        for record in campus_records
        if get_primary_value(record) == 5
    )

    return (
        grade_a_count
        / len(campus_records)
        * 100
    )


# ============================================================
# VARIANCE FORMATTING
# ============================================================

def calculate_variance(record, rule, value):
    method = rule.get("variance_method", "none")

    if method == "none":
        return None

    if record is None:
        return None

    if method == "api_rounded":
        return safe_float(
            record.get("prim_org_prfm_rnd_pct")
        )

    if method == "api_raw":
        return safe_float(
            record.get("prim_org_prfm_raw_pct")
        )

    goal = safe_float(
        rule.get("variance_goal")
    )

    numeric_value = safe_float(value)

    if goal is None or numeric_value is None:
        return None

    variance_sign = rule.get(
        "variance_sign",
        "value_minus_goal",
    )

    if method == "hardcoded_absolute":
        if variance_sign == "goal_minus_value":
            return goal - numeric_value

        return numeric_value - goal

    if method == "hardcoded_relative":
        if goal == 0:
            return None

        if variance_sign == "goal_minus_value":
            return (
                (goal - numeric_value)
                / abs(goal)
                * 100
            )

        return (
            (numeric_value - goal)
            / abs(goal)
            * 100
        )

    return None


def format_variance(variance, variance_format):
    numeric_variance = safe_float(variance)

    if numeric_variance is None:
        return "N/A"

    if variance_format == "percent_0":
        return f"{numeric_variance:,.0f}%"

    if variance_format == "percent_1":
        return f"{numeric_variance:,.1f}%"

    if variance_format == "percent_2":
        return f"{numeric_variance:,.2f}%"

    return str(numeric_variance)


def get_variance_arrow(variance):
    numeric_variance = safe_float(variance)

    if numeric_variance is None:
        return "→"

    if numeric_variance > 0:
        return "↗"

    if numeric_variance < 0:
        return "↘"

    return "→"


def get_variance_color(variance):
    numeric_variance = safe_float(variance)

    if numeric_variance is None:
        return "GRAY"

    if numeric_variance > 0:
        return "GREEN"

    if numeric_variance < 0:
        return "RED"

    return "GREEN"

# ============================================================
# METRIC DRILLDOWN HIERARCHY
# ============================================================

def make_node_key(record):
    metric_id = clean_text(
        record.get("mtrc_id")
    )

    current_level = clean_text(
        record.get("rollup_type")
    )

    current_location = clean_text(
        record.get("location_name")
    )

    return (
        f"{metric_id}|"
        f"{current_level}|"
        f"{current_location}"
    )


def make_parent_key(record):
    metric_id = clean_text(
        record.get("mtrc_id")
    )

    parent_level = clean_text(
        record.get(
            "rollup_identifier_role_lbl"
        )
    )

    parent_location = clean_text(
        record.get("rollup_identifier")
    )

    if not parent_level or not parent_location:
        return ""

    return (
        f"{metric_id}|"
        f"{parent_level}|"
        f"{parent_location}"
    )


def get_rule_for_record(record):
    for rule in METRIC_RULES:
        if rule_matches_record(rule, record):
            return rule

    return None


def get_drilldown_value_format(rule):
    if rule is None:
        return "decimal_1"

    value_format = rule.get(
        "value_format",
        "decimal_1",
    )

    if value_format == "coming_soon":
        return "decimal_1"

    return value_format

def get_display_level(
    metric_name,
    location_type,
    hierarchy_level,
):
    metric_rules = METRIC_DISPLAY_LEVEL_RULES.get(
        metric_name
    )

    # Metrics with confirmed rules must include only
    # the location types explicitly listed.
    if metric_rules is not None:
        return metric_rules.get(
            location_type
        )

    # Default handling for metrics whose exact
    # DataLink drilldown structure is not mapped yet.
    if location_type == "Corporate":
        return 0

    numeric_hierarchy_level = int(
        hierarchy_level or 0
    )

    if numeric_hierarchy_level <= 0:
        return 0

    if numeric_hierarchy_level == 1:
        return 1

    return 2

def create_category_sheet_row(node):
    metric_name = node.get(
        "Metric",
        "",
    )

    location_type = node.get(
        "Current Level",
        "",
    )

    hierarchy_level = node.get(
        "Hierarchy Level",
        0,
    )

    display_level = get_display_level(
        metric_name=metric_name,
        location_type=location_type,
        hierarchy_level=hierarchy_level,
    )

    if display_level is None:
        return None

    return {
        "Metric": metric_name,
        "Level": display_level,
        "Division": (
            node.get("Division")
            or "N/A"
        ),
        "Location Type": location_type,
        "Location": node.get(
            "Current Location",
            "",
        ),
        "Value": node.get(
            "Value",
            "N/A",
        ),
        "Goal": node.get(
            "Goal",
            "N/A",
        ),
        "Variance Direction": node.get(
            "Variance Direction",
            "→",
        ),
        "Variance": node.get(
            "Variance",
            "N/A",
        ),
        "Variance Color": node.get(
            "Variance Color",
            "GRAY",
        ),
        "YTD": node.get(
            "YTD",
            "N/A",
        ),
        "Data As Of": node.get(
            "Data As Of",
            "N/A",
        ),
        "_Category": node.get(
            "Category",
            "",
        ),
    }

def build_category_sheet_rows(
    drilldown_nodes,
):
    rows_by_category = {
        category_code: []
        for category_code in CATEGORY_SHEETS
    }

    metric_order = {
        rule["display_name"]: rule["order"]
        for rule in METRIC_RULES
    }

    seen_rows = set()

    for node in drilldown_nodes:
        category_row = create_category_sheet_row(
            node
        )

        if category_row is None:
            continue


        category_code = category_row.pop(
            "_Category",
            "",
        )

        if category_code not in rows_by_category:
            continue

        row_identity = (
            category_code,
            category_row["Metric"],
            category_row["Level"],
            category_row["Location Type"],
            category_row["Location"],
        )

        if row_identity in seen_rows:
            continue

        seen_rows.add(row_identity)

        rows_by_category[
            category_code
        ].append(category_row)

    for category_code, category_rows in (
        rows_by_category.items()
    ):
        category_rows.sort(
            key=lambda row: (
                metric_order.get(
                    row["Metric"],
                    999,
                ),
                row["Level"],
                row["Location Type"],
                row["Location"],
            )
        )

    return rows_by_category

def get_record_data_as_of(record):
    return (
        clean_text(
            record.get("prim_org_caln_lbl")
        )
        or clean_text(
            record.get("prim_org_time_frame_lbl")
        )
        or "N/A"
    )


def get_record_goal_display(record, rule):
    if rule is None:
        return (
            clean_text(record.get("metric_goal"))
            or "N/A"
        )

    # The top-level hardcoded goal remains authoritative.
    # Lower hierarchy rows may have location-specific goals,
    # so retain the API-provided goal where available.
    api_goal = clean_text(
        record.get("metric_goal")
    )

    if api_goal:
        return api_goal

    return rule.get(
        "goal_display",
        "N/A",
    )


def get_record_variance(record):
    rounded_variance = safe_float(
        record.get("prim_org_prfm_rnd_pct")
    )

    if rounded_variance is not None:
        return rounded_variance

    raw_variance = safe_float(
        record.get("prim_org_prfm_raw_pct")
    )

    return raw_variance


def create_drilldown_node(record):
    rule = get_rule_for_record(record)

    if rule is None:
        return None

    selection = rule.get("selection")

    if selection == "placeholder":
        return None

    node_key = make_node_key(record)
    parent_key = make_parent_key(record)

    if not node_key:
        return None

    metric_id = record.get("mtrc_id")
    variance = get_record_variance(record)

    return {
        "Category": rule["category"],
        "Aspiration": rule["aspiration"],
        "Metric ID": metric_id,
        "Metric": rule["display_name"],
        "Node Key": node_key,
        "Parent Key": parent_key,
        "Hierarchy Level": None,
        "Hierarchy Path": "",
        "Current Level": (
            clean_text(record.get("rollup_type"))
            or "Unknown"
        ),
        "Current Location": (
            clean_text(record.get("location_name"))
            or "Unknown"
        ),
        "Parent Level": clean_text(
            record.get(
                "rollup_identifier_role_lbl"
            )
        ),
        "Parent Location": clean_text(
            record.get("rollup_identifier")
        ),
        "Value": format_value(
            get_primary_value(record),
            get_drilldown_value_format(rule),
        ),
        "Goal": get_record_goal_display(
            record,
            rule,
        ),
        "Variance Direction": get_variance_arrow(
            variance
        ),
        "Variance": format_variance(
            variance,
            rule.get(
                "variance_format",
                "percent_1",
            ),
        ),
        "Variance Color": get_variance_color(
            variance
        ),
        "YTD": format_value(
            get_ytd_value(record),
            rule.get(
                "ytd_format",
                "none",
            ),
        )
        if rule.get("ytd_format") != "none"
        else "N/A",
        "Data As Of": get_record_data_as_of(
            record
        ),
        "Has Children": False,
    }


def build_drilldown_nodes(records):
    nodes_by_key = {}

    for record in records:
        node = create_drilldown_node(record)

        if node is None:
            continue

        node_key = node["Node Key"]

        # Keep one row per unique metric-level-location.
        if node_key not in nodes_by_key:
            nodes_by_key[node_key] = node

    children_by_parent = {}

    for node in nodes_by_key.values():
        parent_key = node["Parent Key"]

        if (
            parent_key
            and parent_key in nodes_by_key
        ):
            children_by_parent.setdefault(
                parent_key,
                [],
            ).append(node)

            nodes_by_key[parent_key][
                "Has Children"
            ] = True

    for children in children_by_parent.values():
        children.sort(
            key=lambda node: (
                node["Current Level"],
                node["Current Location"],
            )
        )

    root_nodes = []

    for node in nodes_by_key.values():
        parent_key = node["Parent Key"]

        if (
            not parent_key
            or parent_key not in nodes_by_key
        ):
            root_nodes.append(node)

    metric_order = {
        rule["display_name"]: rule["order"]
        for rule in METRIC_RULES
    }

    root_nodes.sort(
        key=lambda node: (
            metric_order.get(
                node["Metric"],
                999,
            ),
            node["Current Location"],
        )
    )

    ordered_nodes = []
    visited_keys = set()

    def walk_node(
        node,
        hierarchy_level,
        parent_path,
        inherited_division="",
    ):
        node_key = node["Node Key"]

        if node_key in visited_keys:
            return

        visited_keys.add(node_key)

        current_level = clean_text(
            node.get("Current Level")
        )

        current_location = clean_text(
            node.get("Current Location")
        )

        # If the current record is a Division,
        # it becomes the Division for itself and
        # every child record underneath it.
        if current_level == "Division":
            division_name = current_location

        else:
            division_name = inherited_division

        hierarchy_path = (
            f"{parent_path} > {current_location}"
            if parent_path
            else current_location
        )

        node["Hierarchy Level"] = hierarchy_level
        node["Hierarchy Path"] = hierarchy_path
        node["Division"] = (
            division_name
            or "N/A"
        )

        ordered_nodes.append(node)

        for child_node in children_by_parent.get(
            node_key,
            [],
        ):
            walk_node(
                child_node,
                hierarchy_level + 1,
                hierarchy_path,
                division_name,
            )


        for root_node in root_nodes:
            walk_node(
                root_node,
                0,
                "",
                "",
            )

        # Retain disconnected records instead of losing them.
        for node in nodes_by_key.values():
            if node["Node Key"] not in visited_keys:
                walk_node(
                    node,
                    0,
                    "",
                    "",
                )

    return ordered_nodes

# ============================================================
# BUSINESS TABLE CREATION
# ============================================================

def build_placeholder_row(rule):
    return {
        "Category": rule["category"],
        "Aspiration": rule["aspiration"],
        "Metric": rule["display_name"],
        "Location": "AdventHealth",
        "Value": "Coming Soon",
        "Goal": rule["goal_display"],
        "Variance Direction": "→",
        "Variance": "N/A",
        "Variance Color": "GRAY",
        "YTD": "N/A",
        "Data As Of": "N/A",
    }


def build_standard_business_row(records, rule):
    record = find_corporate_record(
        records,
        rule,
    )

    if record is None:
        return {
            "Category": rule["category"],
            "Aspiration": rule["aspiration"],
            "Metric": rule["display_name"],
            "Location": "AdventHealth",
            "Value": "N/A",
            "Goal": rule["goal_display"],
            "Variance Direction": "→",
            "Variance": "N/A",
            "Variance Color": "GRAY",
            "YTD": "N/A",
            "Data As Of": "N/A",
        }

    value = get_primary_value(record)

    variance = calculate_variance(
        record,
        rule,
        value,
    )

    ytd_format = rule.get(
        "ytd_format",
        "none",
    )

    if ytd_format == "none":
        ytd_display = "N/A"

    else:
        ytd_display = format_value(
            get_ytd_value(record),
            ytd_format,
        )

    data_as_of = (
        clean_text(record.get("prim_org_caln_lbl"))
        or clean_text(
            record.get("prim_org_time_frame_lbl")
        )
        or "N/A"
    )

    return {
        "Category": rule["category"],
        "Aspiration": rule["aspiration"],
        "Metric": rule["display_name"],
        "Location": "AdventHealth",
        "Value": format_value(
            value,
            rule["value_format"],
        ),
        "Goal": rule["goal_display"],
        "Variance Direction": get_variance_arrow(
            variance
        ),
        "Variance": format_variance(
            variance,
            rule["variance_format"],
        ),
        "Variance Color": get_variance_color(
            variance
        ),
        "YTD": ytd_display,
        "Data As Of": data_as_of,
    }


def build_summary_business_row(records, rule):
    selection = rule["selection"]

    if selection == "cms_star_summary":
        summary_value = calculate_cms_star_summary(
            records,
            rule,
        )

        data_as_of = get_latest_metric_period(
            records,
            rule,
        )

    elif selection == "leapfrog_summary":
        summary_value = calculate_leapfrog_summary(
            records,
            rule,
        )

        data_as_of = get_latest_metric_period(
            records,
            rule,
        )

    else:
        summary_value = None
        data_as_of = "N/A"

    return {
        "Category": rule["category"],
        "Aspiration": rule["aspiration"],
        "Metric": rule["display_name"],
        "Location": "AdventHealth",
        "Value": format_value(
            summary_value,
            rule["value_format"],
        ),
        "Goal": rule["goal_display"],
        "Variance Direction": "",
        "Variance": "",
        "Variance Color": "",
        "YTD": "",
        "Data As Of": data_as_of,
    }


def get_latest_metric_period(records, rule):
    matching_records = get_matching_records(
        records,
        rule,
    )

    for record in matching_records:
        period = clean_text(
            record.get("prim_org_caln_lbl")
        )

        if period:
            return period

    return "N/A"


def build_business_table(records):
    business_records = []

    for rule in sorted(
        METRIC_RULES,
        key=lambda item: item["order"],
    ):
        selection = rule["selection"]

        if selection == "placeholder":
            business_record = build_placeholder_row(
                rule
            )

        elif selection in {
            "cms_star_summary",
            "leapfrog_summary",
        }:
            business_record = build_summary_business_row(
                records,
                rule,
            )

        else:
            business_record = build_standard_business_row(
                records,
                rule,
            )

        business_records.append(
            business_record
        )

    if len(business_records) != 20:
        raise ValueError(
            "Expected exactly 20 business metrics, "
            f"but created {len(business_records)}."
        )

    return business_records


def save_business_csv(records):
    business_records = build_business_table(
        records
    )

    columns = [
        "Category",
        "Aspiration",
        "Metric",
        "Location",
        "Value",
        "Goal",
        "Variance Direction",
        "Variance",
        "Variance Color",
        "YTD",
        "Data As Of",
    ]

    with BUSINESS_CSV_OUTPUT_FILE.open(
        mode="w",
        newline="",
        encoding="utf-8-sig",
    ) as csv_file:

        writer = csv.DictWriter(
            csv_file,
            fieldnames=columns,
        )

        writer.writeheader()
        writer.writerows(business_records)

    return business_records

def add_category_sheet(
    workbook,
    category_code,
    category_rows,
):
    sheet_config = CATEGORY_SHEETS[
        category_code
    ]

    worksheet = workbook.create_sheet(
        title=sheet_config["sheet_name"]
    )

    worksheet.append(
        CATEGORY_SHEET_COLUMNS
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    green_font = Font(
        color="006100",
        bold=True,
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    red_font = Font(
        color="9C0006",
        bold=True,
    )

    gray_fill = PatternFill(
        fill_type="solid",
        fgColor="D9E1F2",
    )

    gray_font = Font(
        color="666666",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    column_positions = {
        column_name: column_number
        for column_number, column_name
        in enumerate(
            CATEGORY_SHEET_COLUMNS,
            start=1,
        )
    }

    for category_row in category_rows:
        worksheet.append(
            [
                category_row.get(
                    column_name,
                    "",
                )
                for column_name
                in CATEGORY_SHEET_COLUMNS
            ]
        )

        row_number = worksheet.max_row

        level_cell = worksheet.cell(
            row=row_number,
            column=column_positions["Level"],
        )

        level_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        location_type_cell = worksheet.cell(
            row=row_number,
            column=column_positions[
                "Location Type"
            ],
        )

        location_type_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        if category_row["Level"] == 0:
            for cell in worksheet[row_number]:
                cell.font = Font(
                    bold=True,
                    color="17365D",
                )

        variance_color = clean_text(
            category_row.get(
                "Variance Color"
            )
        ).upper()

        if variance_color == "GREEN":
            selected_fill = green_fill
            selected_font = green_font

        elif variance_color == "RED":
            selected_fill = red_fill
            selected_font = red_font

        else:
            selected_fill = gray_fill
            selected_font = gray_font

        for column_name in [
            "Value",
            "Variance Direction",
            "Variance",
            "Variance Color",
        ]:
            cell = worksheet.cell(
                row=row_number,
                column=column_positions[
                    column_name
                ],
            )

            cell.fill = selected_fill
            cell.font = selected_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    column_widths = {
        "Metric": 38,
        "Level": 10,
        "Division": 58,
        "Location Type": 18,
        "Location": 58,
        "Value": 16,
        "Goal": 18,
        "Variance Direction": 20,
        "Variance": 14,
        "Variance Color": 17,
        "YTD": 12,
        "Data As Of": 15,
    }

    for column_name, width in (
        column_widths.items()
    ):
        worksheet.column_dimensions[
            get_column_letter(
                column_positions[
                    column_name
                ]
            )
        ].width = width

    worksheet.freeze_panes = "F2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )
    worksheet.row_dimensions[1].height = 25

    return worksheet

def add_all_category_sheets(
    workbook,
    drilldown_nodes,
):
    rows_by_category = build_category_sheet_rows(
        drilldown_nodes
    )

    category_row_counts = {}

    for category_code in CATEGORY_SHEETS:
        category_rows = rows_by_category.get(
            category_code,
            [],
        )

        add_category_sheet(
            workbook=workbook,
            category_code=category_code,
            category_rows=category_rows,
        )

        category_row_counts[
            category_code
        ] = len(category_rows)

    return category_row_counts

def add_metric_drilldown_sheet(
    workbook,
    records,
):
    drilldown_nodes = build_drilldown_nodes(
        records
    )

    worksheet = workbook.create_sheet(
        title="Metric Drilldowns"
    )

    columns = [
        "Category",
        "Aspiration",
        "Metric ID",
        "Metric",
        "Node Key",
        "Parent Key",
        "Hierarchy Level",
        "Hierarchy Path",
        "Current Level",
        "Current Location",
        "Parent Level",
        "Parent Location",
        "Value",
        "Goal",
        "Variance Direction",
        "Variance",
        "Variance Color",
        "YTD",
        "Data As Of",
    ]

    worksheet.append(columns)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    green_font = Font(
        color="006100",
        bold=True,
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    red_font = Font(
        color="9C0006",
        bold=True,
    )

    gray_fill = PatternFill(
        fill_type="solid",
        fgColor="D9E1F2",
    )

    gray_font = Font(
        color="666666",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    column_positions = {
        column_name: column_number
        for column_number, column_name
        in enumerate(columns, start=1)
    }

    row_number_by_node_key = {}

    for drilldown_node in drilldown_nodes:
        worksheet.append(
            [
                drilldown_node.get(
                    column,
                    "",
                )
                for column in columns
            ]
        )

        row_number = worksheet.max_row

        row_number_by_node_key[
            drilldown_node["Node Key"]
        ] = row_number

        hierarchy_level = int(
            drilldown_node.get(
                "Hierarchy Level",
                0,
            )
            or 0
        )

        # Excel supports outline levels 0 through 7
        # through openpyxl, representing 8 visible levels.
        worksheet.row_dimensions[
            row_number
        ].outlineLevel = min(
            hierarchy_level,
            7,
        )

        # Start with all child levels collapsed.
        worksheet.row_dimensions[
            row_number
        ].hidden = hierarchy_level > 0

        if drilldown_node.get(
            "Has Children"
        ):
            worksheet.row_dimensions[
                row_number
            ].collapsed = True

        location_cell = worksheet.cell(
            row=row_number,
            column=column_positions[
                "Current Location"
            ],
        )

        location_cell.alignment = Alignment(
            vertical="center",
            indent=min(
                hierarchy_level,
                15,
            ),
        )

        if hierarchy_level == 0:
            location_cell.font = Font(
                bold=True,
                color="17365D",
            )

        variance_color = clean_text(
            drilldown_node.get(
                "Variance Color"
            )
        ).upper()

        if variance_color == "GREEN":
            selected_fill = green_fill
            selected_font = green_font

        elif variance_color == "RED":
            selected_fill = red_fill
            selected_font = red_font

        else:
            selected_fill = gray_fill
            selected_font = gray_font

        for column_name in [
            "Value",
            "Variance Direction",
            "Variance",
            "Variance Color",
        ]:
            cell = worksheet.cell(
                row=row_number,
                column=column_positions[
                    column_name
                ],
            )

            cell.fill = selected_fill
            cell.font = selected_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    worksheet.sheet_properties.outlinePr.summaryBelow = (
        False
    )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    column_widths = {
        "Category": 13,
        "Aspiration": 38,
        "Metric ID": 12,
        "Metric": 34,
        "Node Key": 65,
        "Parent Key": 65,
        "Hierarchy Level": 16,
        "Hierarchy Path": 90,
        "Current Level": 18,
        "Current Location": 58,
        "Parent Level": 18,
        "Parent Location": 58,
        "Value": 16,
        "Goal": 18,
        "Variance Direction": 20,
        "Variance": 14,
        "Variance Color": 17,
        "YTD": 12,
        "Data As Of": 15,
    }

    for column_name, width in column_widths.items():
        worksheet.column_dimensions[
            get_column_letter(
                column_positions[column_name]
            )
        ].width = width

    # Hide technical relationship columns while retaining
    # them in the workbook for QA and reconstruction.
    for technical_column in [
        "Metric ID",
        "Node Key",
        "Parent Key",
        "Hierarchy Level",
    ]:
        worksheet.column_dimensions[
            get_column_letter(
                column_positions[
                    technical_column
                ]
            )
        ].hidden = True

    return drilldown_nodes


def save_business_excel(
    business_records,
    records,):
    columns = [
        "Category",
        "Aspiration",
        "Metric",
        "Location",
        "Value",
        "Goal",
        "Variance Direction",
        "Variance",
        "Variance Color",
        "YTD",
        "Data As Of",
    ]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Vision 2030"

    # Header formatting
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    # Status colors
    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    green_font = Font(
        color="006100",
        bold=True,
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    red_font = Font(
        color="9C0006",
        bold=True,
    )

    gray_fill = PatternFill(
        fill_type="solid",
        fgColor="D9E1F2",
    )

    gray_font = Font(
        color="666666",
        bold=True,
    )

    # Write header
    worksheet.append(columns)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Write business records
    for business_record in business_records:
        worksheet.append(
            [
                business_record.get(column, "")
                for column in columns
            ]
        )

    column_positions = {
        column_name: column_number
        for column_number, column_name
        in enumerate(columns, start=1)
    }

    # Columns that will receive the corresponding status color
    columns_to_color = [
        "Value",
        "Variance Direction",
        "Variance",
        "Variance Color",
    ]

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        variance_color = str(
            worksheet.cell(
                row=row_number,
                column=column_positions["Variance Color"],
            ).value
            or ""
        ).strip().upper()

        if variance_color == "GREEN":
            selected_fill = green_fill
            selected_font = green_font

        elif variance_color == "RED":
            selected_fill = red_fill
            selected_font = red_font

        else:
            selected_fill = gray_fill
            selected_font = gray_font

        for column_name in columns_to_color:
            cell = worksheet.cell(
                row=row_number,
                column=column_positions[column_name],
            )

            cell.fill = selected_fill
            cell.font = selected_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # Align all remaining cells
    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
            )

    # Column widths
    column_widths = {
        "Category": 13,
        "Aspiration": 38,
        "Metric": 38,
        "Location": 18,
        "Value": 16,
        "Goal": 18,
        "Variance Direction": 20,
        "Variance": 14,
        "Variance Color": 17,
        "YTD": 12,
        "Data As Of": 15,
    }

    for column_name, width in column_widths.items():
        column_number = column_positions[column_name]

        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 25

    drilldown_nodes = build_drilldown_nodes(
        records
    )

    category_row_counts = add_all_category_sheets(
        workbook=workbook,
        drilldown_nodes=drilldown_nodes,
    )

    workbook.save(
        BUSINESS_EXCEL_OUTPUT_FILE
    )

    return drilldown_nodes

# ============================================================
# MAIN PROCESS
# ============================================================

def main():
    OUTPUT_FOLDER.mkdir(
        parents=True,
        exist_ok=True,
    )

    print(
        "\nStarting Vision 2030 DataLink extraction..."
    )
    print(
        f"Python interpreter: {sys.executable}"
    )
    print(
        f"Output folder: {OUTPUT_FOLDER}"
    )

    edge_options = Options()
    edge_options.add_argument(
        "--start-maximized"
    )

    driver = None

    try:
        print("\nOpening Microsoft Edge...")

        driver = webdriver.Edge(
            options=edge_options
        )

        print(
            "Opening the Vision 2030 DataLink page..."
        )

        driver.get(
            DATALINK_PAGE_URL
        )

        input(
            "\nComplete the normal DataLink sign-in in Edge "
            "if prompted.\n"
            "After the Vision 2030 page is fully loaded, "
            "return here and press Enter..."
        )

        print(
            "\nOpening the Vision 2030 API endpoint..."
        )

        driver.get(
            API_URL
        )

        WebDriverWait(
            driver,
            120,
        ).until(
            lambda browser: bool(
                get_json_text_from_browser(
                    browser
                )
            )
        )

        response_text = (
            get_json_text_from_browser(driver)
        )

        if not response_text:
            raise ValueError(
                "The browser page did not contain "
                "an API response."
            )

        print(
            "API response received."
        )

        try:
            records = json.loads(
                response_text
            )

        except json.JSONDecodeError as error:
            debug_file = (
                OUTPUT_FOLDER
                / (
                    "vision2030_unparsed_response_"
                    f"{TIMESTAMP}.txt"
                )
            )

            debug_file.write_text(
                response_text,
                encoding="utf-8",
            )

            raise ValueError(
                "The API response could not be parsed "
                "as JSON. The unparsed response was "
                f"saved to: {debug_file}"
            ) from error

        validate_records(records)

        save_json(records)

        raw_columns = save_raw_csv(
            records
        )

        business_records = build_business_table(
            records
        )

        drilldown_nodes = save_business_excel(
            business_records,
            records,
        )

        unique_metrics = {
            record.get("metric_name")
            for record in records
            if record.get("metric_name")
        }

        unique_locations = {
            record.get("location_name")
            for record in records
            if record.get("location_name")
        }

        unique_rollup_types = {
            record.get("rollup_type")
            for record in records
            if record.get("rollup_type")
        }

        missing_business_metrics = [
            record["Metric"]
            for record in business_records
            if record["Value"] == "N/A"
        ]

        print(
            "\n========================================"
        )
        print(
            "VISION 2030 EXTRACTION PASSED"
        )
        print(
            "========================================"
        )
        print(
            f"API records extracted: {len(records):,}"
        )
        print(
            f"Raw CSV columns: {len(raw_columns):,}"
        )
        print(
            f"Unique API metrics: {len(unique_metrics):,}"
        )
        print(
            f"Unique locations: {len(unique_locations):,}"
        )
        print(
            "Rollup types: "
            + ", ".join(
                sorted(unique_rollup_types)
            )
        )
        print(
            f"Business metrics created: "
            f"{len(business_records)}"
        )

        print(
            f"Metric drilldown rows created: "
            f"{len(drilldown_nodes):,}"
        )

        if missing_business_metrics:
            print(
                "Business metrics with no matching "
                "corporate API value:"
            )

            for metric_name in missing_business_metrics:
                print(
                    f"  - {metric_name}"
                )

        print(
            "----------------------------------------"
        )
        print(
            f"Raw JSON:\n{JSON_OUTPUT_FILE}"
        )
        print(
            f"\nRaw API CSV:\n{RAW_CSV_OUTPUT_FILE}"
        )
        print(
            "\nColor-formatted Excel workbook:\n"
            f"{BUSINESS_EXCEL_OUTPUT_FILE}"
        )
        print(
            "\nColor-formatted Excel workbook:\n"
            f"{BUSINESS_EXCEL_OUTPUT_FILE}"
        )
        print(
            "========================================"
        )

    except Exception as error:
        print(
            "\n========================================"
        )
        print(
            "VISION 2030 EXTRACTION FAILED"
        )
        print(
            "========================================"
        )
        print(
            f"{type(error).__name__}: {error}"
        )
        print(
            "========================================"
        )

        raise

    finally:
        if driver is not None:
            input(
                "\nPress Enter to close "
                "the Edge window..."
            )

            driver.quit()


if __name__ == "__main__":
    main()