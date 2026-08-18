from __future__ import annotations

import argparse
import json
import math
import re
import threading
import webbrowser
from collections import Counter, defaultdict
from datetime import datetime
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HOST = "127.0.0.1"
DEFAULT_PORT = 8000
SITE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SITE_DIR.parent
OUTPUT_ROOT = PROJECT_DIR / "output"
GENERATED_DIR = SITE_DIR / "generated"
WORKBOOK_PATTERN = "vision2030_business_table_*.xlsx"
SUMMARY_SHEET = "Vision 2030"
DETAIL_SHEETS = ("Learning", "Team", "Consumer", "Clinical", "Financial", "Risk", "WPC")
VALID_STATUS = {"GREEN", "RED", "GRAY"}
MISSING_TEXT = {"", "N/A", "NA", "NONE", "NULL", "-", "COMING SOON"}


def clean(value: Any) -> str:
    if value is None:
        return "N/A"
    text = str(value).strip()
    return text if text else "N/A"


def parse_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip().upper().replace(",", "")
    if text in MISSING_TEXT:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace("$", "").replace("%", "")
    match = re.search(r"[-+]?\d*\.?\d+", text)
    if not match:
        return None
    number = float(match.group())
    if "B" in text:
        number *= 1_000_000_000
    elif "M" in text:
        number *= 1_000_000
    elif "K" in text:
        number *= 1_000
    return -number if negative else number


def normalize_status(value: Any) -> str:
    status = clean(value).upper()
    return status if status in VALID_STATUS else "GRAY"


def row_dict(headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
    return {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}


def find_latest_workbook(output_root: Path) -> Path:
    if not output_root.exists():
        raise FileNotFoundError(f"Output folder does not exist: {output_root}")
    candidates = [p for p in output_root.glob(f"*/{WORKBOOK_PATTERN}") if p.is_file() and not p.name.startswith("~$")]
    if not candidates:
        candidates = [p for p in output_root.rglob(WORKBOOK_PATTERN) if p.is_file() and not p.name.startswith("~$")]
    if not candidates:
        raise FileNotFoundError(f"No {WORKBOOK_PATTERN} found under {output_root}")
    timestamp_pattern = re.compile(r"^\d{4}-\d{2}-\d{2}_\d{6}$")
    timestamped = [p for p in candidates if timestamp_pattern.match(p.parent.name)]
    if timestamped:
        return max(timestamped, key=lambda p: (p.parent.name, p.name))
    return max(candidates, key=lambda p: (p.stat().st_mtime, p.name))


def read_sheet(ws) -> list[dict[str, Any]]:
    iterator = ws.iter_rows(values_only=True)
    raw_headers = next(iterator)
    headers = [clean(value) for value in raw_headers]
    rows = []
    for values in iterator:
        if not any(value is not None and str(value).strip() for value in values):
            continue
        rows.append(row_dict(headers, values))
    return rows


def metric_record(row: dict[str, Any], category: str, detail: bool) -> dict[str, Any]:
    record = {
        "category": category.upper(),
        "metric": clean(row.get("Metric")),
        "division": clean(row.get("Division")) if detail else "N/A",
        "location_type": clean(row.get("Location Type")) if detail else "Corporate",
        "location": clean(row.get("Location")),
        "value": clean(row.get("Value")),
        "value_numeric": parse_number(row.get("Value")),
        "goal": clean(row.get("Goal")),
        "variance_direction": clean(row.get("Variance Direction")),
        "variance": clean(row.get("Variance")),
        "status": normalize_status(row.get("Variance Color")),
        "ytd": clean(row.get("YTD")),
        "data_as_of": clean(row.get("Data As Of")),
    }
    if detail:
        level = row.get("Level")
        try:
            record["level"] = int(level) if level is not None else None
        except (TypeError, ValueError):
            record["level"] = None
    else:
        record["aspiration"] = clean(row.get("Aspiration"))
    return record


def calculate_status_counts(records: list[dict[str, Any]]) -> dict[str, int]:
    counts = Counter(record["status"] for record in records)
    return {status.lower(): counts.get(status, 0) for status in ("GREEN", "RED", "GRAY")}


def calculate_achievement(records: list[dict[str, Any]]) -> dict[str, Any]:
    counts = calculate_status_counts(records)
    eligible = counts["green"] + counts["red"]
    percent = round((counts["green"] / eligible) * 100, 1) if eligible else None
    return {"counts": counts, "eligible": eligible, "percent": percent}


def build_category_summary(summary_records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in summary_records:
        grouped[record["category"]].append(record)
    result = []
    for category, records in grouped.items():
        result.append({
            "category": category,
            "aspiration": records[0].get("aspiration", "N/A"),
            "metrics": len(records),
            **calculate_achievement(records),
        })
    return result


def build_division_comparison(detail_records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    allowed = {"CFD", "EFD", "WFD", "PHD", "MSD", "MSD 1", "MSD 2"}
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in detail_records:
        division = record["division"]
        if division in allowed and record["status"] in {"GREEN", "RED", "GRAY"}:
            grouped[division].append(record)
    output = []
    for division in sorted(grouped):
        records = grouped[division]
        unique = {}
        for record in records:
            key = (record["metric"], record["location_type"], record["location"], record["data_as_of"])
            unique[key] = record
        deduped = list(unique.values())
        output.append({"division": division, "records": len(deduped), **calculate_achievement(deduped)})
    return output


def build_rankings(detail_records: list[dict[str, Any]], limit: int = 10) -> dict[str, Any]:
    by_metric: dict[str, list[dict[str, Any]]] = defaultdict(list)
    permitted_types = {"DIVISION", "REGION", "MARKET", "FACILITY", "HOSPITAL", "CAMPUS", "BUSINESS UNIT"}
    for record in detail_records:
        if record["value_numeric"] is None:
            continue
        if record["location_type"].upper() not in permitted_types:
            continue
        by_metric[record["metric"]].append(record)
    result = {}
    for metric, records in by_metric.items():
        deduped = {}
        for record in records:
            key = (record["location_type"], record["location"])
            deduped[key] = record
        ordered = sorted(deduped.values(), key=lambda item: item["value_numeric"])
        result[metric] = {"lowest": ordered[:limit], "highest": list(reversed(ordered[-limit:]))}
    return result


def build_data(workbook_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    missing = [name for name in (SUMMARY_SHEET, *DETAIL_SHEETS) if name not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Workbook is missing required sheets: {', '.join(missing)}")
    summary_rows = read_sheet(workbook[SUMMARY_SHEET])
    summary_records = [metric_record(row, clean(row.get("Category")), False) for row in summary_rows]
    detail_records = []
    sheet_counts = {}
    for sheet_name in DETAIL_SHEETS:
        rows = read_sheet(workbook[sheet_name])
        sheet_counts[sheet_name] = len(rows)
        detail_records.extend(metric_record(row, sheet_name, True) for row in rows)
    metrics = sorted({record["metric"] for record in detail_records if record["metric"] != "N/A"})
    location_types = sorted({record["location_type"] for record in detail_records if record["location_type"] != "N/A"})
    divisions = sorted({record["division"] for record in detail_records if record["division"] != "N/A"})
    return {
        "metadata": {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "source_workbook": str(workbook_path),
            "source_folder": str(workbook_path.parent),
            "sheet_rows": sheet_counts,
            "summary_metrics": len(summary_records),
            "detail_rows": len(detail_records),
            "unique_metrics": len(metrics),
        },
        "summary": summary_records,
        "overall": calculate_achievement(summary_records),
        "categories": build_category_summary(summary_records),
        "division_comparison": build_division_comparison(detail_records),
        "rankings": build_rankings(detail_records),
        "filters": {"metrics": metrics, "location_types": location_types, "divisions": divisions},
        "details": detail_records,
    }


def write_data(data: dict[str, Any]) -> Path:
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    output_path = GENERATED_DIR / "website_data.json"
    temporary_path = output_path.with_suffix(".json.tmp")
    temporary_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary_path.replace(output_path)
    return output_path


class NoCacheHandler(SimpleHTTPRequestHandler):
    def end_headers(self) -> None:
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        super().end_headers()


def serve(port: int) -> None:
    import os
    os.chdir(SITE_DIR)
    server = None
    selected_port = port
    for candidate_port in range(port, port + 11):
        try:
            server = ThreadingHTTPServer((HOST, candidate_port), NoCacheHandler)
            selected_port = candidate_port
            break
        except OSError:
            continue
    if server is None:
        raise OSError(f"No available local port from {port} through {port + 10}.")
    url = f"http://{HOST}:{selected_port}/briefing.html"
    print(f"Website: {url}")
    print("Press Ctrl+C to stop.")
    threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nServer stopped.")
    finally:
        server.server_close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Build and launch the Vision 2030 factual website.")
    parser.add_argument("--output-root", type=Path, default=OUTPUT_ROOT)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--port", type=int, default=DEFAULT_PORT)
    parser.add_argument("--build-only", action="store_true")
    args = parser.parse_args()
    workbook_path = args.workbook.resolve() if args.workbook else find_latest_workbook(args.output_root.resolve())
    print(f"Latest workbook selected: {workbook_path}")
    data = build_data(workbook_path)
    output_path = write_data(data)
    print(f"Website data generated: {output_path}")
    print(f"Summary metrics: {data['metadata']['summary_metrics']}")
    print(f"Detail rows: {data['metadata']['detail_rows']}")
    print(f"Unique metrics: {data['metadata']['unique_metrics']}")
    if not args.build_only:
        serve(args.port)


if __name__ == "__main__":
    main()
