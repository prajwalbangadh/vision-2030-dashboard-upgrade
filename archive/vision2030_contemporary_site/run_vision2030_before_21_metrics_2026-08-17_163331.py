from __future__ import annotations
import argparse,csv,json,os,re,threading,webbrowser
from collections import Counter,defaultdict
from copy import deepcopy
from datetime import datetime
from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

SITE=Path(__file__).resolve().parent; PROJECT=SITE.parent; OUTPUT=PROJECT/'output'; TEMPLATES=SITE/'templates'; GENERATED=SITE/'generated'
PAGES=('briefing.html','insights.html','data_explorer.html')
SHEETS=('Learning','Team','Consumer','Clinical','Financial','Risk','WPC')
ALIASES={'Internal Leader Fill Rate':'Internal Leader Fill Rate','Leader Effectiveness':'Leader Effectiveness (Facility Leadership)','Opportunity to Learn and Grow':'Opportunity to Learn and Grow','Forecast and Hire':'Forecast and Hire','Team Member Engagement':'Team Member Engagement','Total Turnover (Rolling 12)':'Total Turnover','Physician Engagement':'Physician Engagement','Annual Active Users':'Digital Tool Utilization','LTR ED':'Patient Experience – ED LTR','LTR Inpatient':'Patient Experience – Inpatient LTR','LTR Med Practice':'Patient Experience – Med Practice LTR','CMS Star Rating':'CMS Overall Hospital Star Rating','Leapfrog Safety Grade':'Leapfrog Hospital Safety Grade','All-Adult Inpatient Mortality':'All-Adult Inpatient Mortality','Length of Stay O/E':'Length of Stay O/E','EBITDA Dollars':'EBITDA $','TOR Dollars':'TOR $','TOR Growth Rate':'TOR Growth Rate','Domestic Spend':'Domestic Spend'}
PILLAR={'Learning':'Dynamic Learning Community','Team':'Team Member Promise','Consumer':'Consumer Focused Connected Network','Clinical':'Clinical Excellence','Financial':'Financial Strength & Growth','Risk':'Managed Population Risk','WPC':'Whole-Person Care'}
DIV_LABEL={'CFD':'CFD','EFD':'EFD','WFD':'WFD','PHD':'PHD','MSD':'MSD','MSD 1':'MSD 1','MSD 2':'MSD 2','Corporate Services Division Cost Centers by Reporting Hierarchy':'Corporate Services','Florida Division Cost Centers by Reporting Hierarchy':'Florida Division (Unresolved)','Multistate Division Cost Centers by Reporting Hierarchy':'Multistate Division (Unresolved)'}
DIV_ORDER=['CFD','EFD','WFD','PHD','MSD','MSD 1','MSD 2','Corporate Services Division Cost Centers by Reporting Hierarchy','Florida Division Cost Centers by Reporting Hierarchy','Multistate Division Cost Centers by Reporting Hierarchy']
VALID={'green','red','gray'}

def text(v): return 'N/A' if v is None or not str(v).strip() else str(v).strip()
def status(v):
 x=text(v).lower(); return x if x in VALID else 'gray'
def slug(v): return re.sub(r'[^A-Z0-9]+','_',text(v).upper()).strip('_') or 'NA'
def latest(root):
 files=[p for p in root.glob('*/vision2030_business_table_*.xlsx') if p.is_file() and not p.name.startswith('~$')]
 if not files: raise FileNotFoundError(f'No business workbook found under {root}')
 stamped=[p for p in files if re.fullmatch(r'\d{4}-\d{2}-\d{2}_\d{6}',p.parent.name)]
 return max(stamped or files,key=lambda p:(p.parent.name,p.name))
def rows(ws):
 it=ws.iter_rows(values_only=True); hdr=[text(v) for v in next(it)]
 for n,vals in enumerate(it,2):
  if any(v is not None and str(v).strip() for v in vals):
   d={hdr[i]:vals[i] if i<len(vals) else None for i in range(len(hdr))}; d['_row']=n; yield d
def model(html):
 marker='const VD ='; start=html.index(marker)+len(marker); value,used=json.JSONDecoder().raw_decode(html[start:]); return value,start,start+used
def scope_id(location_type,location,division):
 lt=text(location_type); loc=text(location); div=text(division)
 if lt.lower()=='corporate' and loc.casefold()=='adventhealth': return 'SYS'
 if lt.lower()=='division':
  for key,label in DIV_LABEL.items():
   if loc.casefold() in {key.casefold(),label.casefold()} or div.casefold()==key.casefold():
    return key if key in {'CFD','EFD','WFD','PHD','MSD','MSD 1','MSD 2'} else 'DIV::'+slug(key)
 return f'{slug(lt)}::{slug(loc)}'

def friendly_scope_label(location_type,location,division):
 lt=text(location_type); loc=text(location); div=text(division)
 # These are distinct Region-level rollups in the report, so keep the data
 # separate but display normalized Division terminology everywhere.
 if lt.lower()=='region' and div in {'CFD','EFD','WFD','PHD','MSD','MSD 1','MSD 2'}:
  hierarchy_tokens=('division cost centers by reporting hierarchy','division - region','primary health cost centers by reporting hierarchy')
  if any(token in loc.casefold() for token in hierarchy_tokens):
   return f'{div} Reporting Hierarchy'
 replacements={
  'Central Florida Division':'CFD','East Florida Division':'EFD','West Florida Division':'WFD',
  'Primary Health Division':'PHD','Multi-State Division':'MSD','Multistate Division':'MSD',
 }
 for long_name,short_name in replacements.items():
  loc=re.sub(re.escape(long_name),short_name,loc,flags=re.IGNORECASE)
 loc=re.sub(r'\s+Cost Centers by Reporting Hierarchy$','',loc,flags=re.IGNORECASE)
 return loc

def selector_group(location_type):
    location_type_text = text(location_type)
    normalized_type = location_type_text.casefold()

    if normalized_type == "corporate":
        return "System"

    if normalized_type == "division":
        return "Division Summary"

    if normalized_type == "region":
        return "Regions"

    if normalized_type == "market":
        return "Markets"

    if normalized_type == "facility":
        return "Facilities"

    return "Other"


def selector_subgroup(location_type):
    location_type_text = text(location_type)
    normalized_type = location_type_text.casefold()

    subgroup_labels = {
        "business unit": "Business Units",
        "campus": "Campuses",
        "cost center": "Cost Centers",
        "costing company": "Costing Companies",
        "frl": "FRLs",
        "hospital": "Hospitals",
    }

    return subgroup_labels.get(
        normalized_type,
        location_type_text,
    )


def normalized_division_key(division):
    division_text = text(division)

    confirmed_divisions = {
        "CFD",
        "EFD",
        "WFD",
        "PHD",
        "MSD",
        "MSD 1",
        "MSD 2",
    }

    if division_text in confirmed_divisions:
        return division_text

    if division_text == "Corporate Services":
        return "Corporate Services"

    if division_text in {
        "Corporate Services Division Cost Centers by Reporting Hierarchy",
        "Corporate Services Division",
    }:
        return "Corporate Services"

    if division_text in {
        "Florida Division Cost Centers by Reporting Hierarchy",
        "Florida Division",
    }:
        return "Florida Division (Unresolved)"

    if division_text in {
        "Multistate Division Cost Centers by Reporting Hierarchy",
        "Multi-State Division Cost Centers by Reporting Hierarchy",
        "Multistate Division",
        "Multi-State Division",
    }:
        return "Multistate Division (Unresolved)"

    return "Unclassified"


def classify(lt):
 x=text(lt).lower()
 if x=='corporate': return 'system'
 if x=='division': return 'division'
 if x in {'region','market'}: return 'region'
 return 'facility'
def parent_id(lt,division):
 x=text(lt).lower(); div=text(division)
 if x=='corporate': return None
 if x=='division': return 'SYS'
 if div in {'MSD 1','MSD 2'}: return div
 if div in {'CFD','EFD','WFD','PHD','MSD'}: return div
 if div in DIV_LABEL: return 'DIV::'+slug(div)
 return 'UNCLASS::'+slug(text(lt))
def extract(book):
 summary=list(rows(book['Vision 2030'])); details=[]
 for sh in SHEETS:
  for r in rows(book[sh]): r['_sheet']=sh; details.append(r)
 return summary,details
def build(template,book,workbook_path):
 summary,details=extract(book); old_metric_by_name={m['metric']:m for m in template['metrics']}
 metrics=[]; metric_names=[]
 for r in summary:
  src=text(r.get('Metric')); ui=ALIASES.get(src)
  if not ui: continue
  metric_names.append(ui); base=deepcopy(old_metric_by_name.get(ui,{})); base.update({'metric':ui,'pillar':text(r.get('Aspiration')),'goal':text(r.get('Goal'))}); metrics.append(base)
 # Add source metrics not present in summary only if aliased and not already present.
 for sh in SHEETS:
  for src,ui in ALIASES.items():
   if ui not in metric_names and any(text(r.get('Metric'))==src and r['_sheet']==sh for r in details):
    sample=next(r for r in details if text(r.get('Metric'))==src and r['_sheet']==sh); metrics.append({'metric':ui,'pillar':PILLAR[sh],'goal':text(sample.get('Goal'))}); metric_names.append(ui)
 scopes=[{
  'id':'SYS',
  'label':'AdventHealth System',
  'kind':'system',
  'parent':None,
  'asp':0,
  'divisionKey':'SYS',
  'selectorGroup':'System',
  'selectorSubgroup':'System'
 }]; seen={'SYS'}
 # generated division parents
 observed_div={text(r.get('Division')) for r in details if text(r.get('Division'))!='N/A'}
 for div in DIV_ORDER:
  if div in observed_div or div in {'CFD','EFD','WFD','PHD','MSD','MSD 1','MSD 2'}:
   sid=div if div in {'CFD','EFD','WFD','PHD','MSD','MSD 1','MSD 2'} else 'DIV::'+slug(div)
   parent='MSD' if div in {'MSD 1','MSD 2'} else 'SYS'
   scopes.append({
    'id':sid,
    'label':DIV_LABEL[div],
    'kind':'division',
    'parent':parent,
    'asp':0,
    'divisionKey':normalized_division_key(div),
    'selectorGroup':'Division Summary',
    'selectorSubgroup':'Division Summary'
   })
   seen.add(sid)
 # unclassified grouping parents
 un_types=sorted({text(r.get('Location Type')) for r in details if text(r.get('Division')) not in observed_div or text(r.get('Division'))=='N/A'})
 scopes.append({
  'id':'UNCLASS',
  'label':'Unclassified / Other',
  'kind':'division',
  'parent':'SYS',
  'asp':0,
  'divisionKey':'Unclassified',
  'selectorGroup':'Division Summary',
  'selectorSubgroup':'Unclassified'
 })
 seen.add('UNCLASS')
 all_types=sorted({text(r.get('Location Type')) for r in details})
 for lt in all_types:
  sid='UNCLASS::'+slug(lt); scopes.append({
   'id':sid,
   'label':f'Unclassified {lt}',
   'kind':'region',
   'parent':'UNCLASS',
   'asp':0,
   'divisionKey':'Unclassified',
   'selectorGroup':selector_group(lt),
   'selectorSubgroup':selector_subgroup(lt)
  })
  seen.add(sid)
 data=defaultdict(lambda:{'values':{},'scores':{}}); source_map={}; duplicates=[]
 # system
 for r in summary:
  ui=ALIASES.get(text(r.get('Metric')))
  if ui:
   data['SYS']['values'][ui]=text(r.get('Value')); data['SYS']['scores'][ui]=status(r.get('Variance Color')); source_map[('SYS',ui)]=('Vision 2030',r['_row'],r)
 # detail scopes
 for r in details:
  ui=ALIASES.get(text(r.get('Metric')))
  if not ui: continue
  sid=scope_id(r.get('Location Type'),r.get('Location'),r.get('Division'))
  if sid not in seen:
   parent=parent_id(r.get('Location Type'),r.get('Division'))
   if parent not in seen: parent='UNCLASS::'+slug(r.get('Location Type'))
   scopes.append({
    'id':sid,
    'label':friendly_scope_label(
     r.get('Location Type'),
     r.get('Location'),
     r.get('Division')
    ),
    'kind':classify(r.get('Location Type')),
    'parent':parent,
    'asp':0,
    'sourceType':text(r.get('Location Type')),
    'divisionKey':normalized_division_key(r.get('Division')),
    'selectorGroup':selector_group(r.get('Location Type')),
    'selectorSubgroup':selector_subgroup(r.get('Location Type'))
   })
   seen.add(sid)
  key=(sid,ui)
  # The Vision 2030 summary sheet is authoritative for SYS. Detail-sheet Corporate rows
  # repeat the same value but may carry GRAY because detail color is not populated.
  if sid=='SYS' and key in source_map and source_map[key][0]=='Vision 2030':
   continue
  if key in source_map:
   prior=source_map[key][2]
   if text(prior.get('Value'))!=text(r.get('Value')) or status(prior.get('Variance Color'))!=status(r.get('Variance Color')): duplicates.append((sid,ui,source_map[key],(r['_sheet'],r['_row'],r)))
   continue
  data[sid]['values'][ui]=text(r.get('Value')); data[sid]['scores'][ui]=status(r.get('Variance Color')); source_map[key]=(r['_sheet'],r['_row'],r)
 facilities=[]
 for sc in scopes:
  if sc['kind']=='facility': facilities.append({'name':sc['label'],'div':sc['parent'],'division':next((x['label'] for x in scopes if x['id']==sc['parent']),sc['parent']),'region':None,'ceo':'N/A','start':'N/A','asp':0,'short':sc['label']})
 generated=deepcopy(template); generated.update({'period':max([text(r.get('Data As Of')) for r in summary if text(r.get('Data As Of'))!='N/A'],key=lambda x:sum(text(z.get('Data As Of'))==x for z in summary)),'metrics':metrics,'scopes':scopes,'metricData':dict(data),'facilities':facilities,'aspiration':{},'regionMeta':{},'divName':{s['id']:s['label'] for s in scopes if s['kind']=='division'},'_source':{'workbook':str(workbook_path),'builtAt':datetime.now().isoformat(timespec='seconds')}})
 return generated,summary,details,source_map,duplicates

def audit(generated,summary,details,source_map,duplicates):
 rows_out=[]; mismatches=0
 for (sid,ui),(sheet,rownum,r) in source_map.items():
  sv=generated['metricData'][sid]['values'][ui]; ss=generated['metricData'][sid]['scores'][ui]; wv=text(r.get('Value')); ws=status(r.get('Variance Color'))
  result='PASS' if sv==wv and ss==ws else 'FAIL'; mismatches += result!='PASS'
  rows_out.append({'Scope ID':sid,'UI Metric':ui,'Source Sheet':sheet,'Source Row':rownum,'Location Type':text(r.get('Location Type')),'Location':text(r.get('Location')),'Division':text(r.get('Division')),'Website Value':sv,'Workbook Value':wv,'Website Status':ss,'Workbook Status':ws,'Goal':text(r.get('Goal')),'Data As Of':text(r.get('Data As Of')),'Result':result})
 GENERATED.mkdir(exist_ok=True)
 with (GENERATED/'scope_metric_audit.csv').open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.DictWriter(f,fieldnames=list(rows_out[0])); w.writeheader(); w.writerows(rows_out)
 summary_json={'sourceRowsMapped':len(source_map),'generatedScopes':len(generated['scopes']),'generatedMetrics':len(generated['metrics']),'valueStatusMismatches':mismatches,'conflictingDuplicateKeys':len(duplicates),'auditPass':mismatches==0 and len(duplicates)==0}
 (GENERATED/'validation_summary.json').write_text(json.dumps(summary_json,indent=2),encoding='utf-8')
 return summary_json

def inject(template_path,out_path,generated):
 html=template_path.read_text(encoding='utf-8'); old,start,end=model(html); payload=json.dumps(generated,ensure_ascii=False,separators=(',',':')); rendered=html[:start]+payload+html[end:]; temp=out_path.with_suffix('.tmp'); temp.write_text(rendered,encoding='utf-8'); temp.replace(out_path); check,cs,ce=model(rendered); assert check['_source']==generated['_source']; assert html[:start]==rendered[:cs] and html[end:]==rendered[ce:]
class Handler(SimpleHTTPRequestHandler):
 def end_headers(self): self.send_header('Cache-Control','no-store'); super().end_headers()
def serve(port):
 os.chdir(SITE); server=None
 for p in range(port,port+11):
  try: server=ThreadingHTTPServer(('127.0.0.1',p),Handler); port=p; break
  except OSError: pass
 if server is None: raise OSError('No available port')
 url=f'http://127.0.0.1:{port}/briefing.html'; print('Website:',url); threading.Timer(.7,lambda:webbrowser.open(url)).start()
 try: server.serve_forever()
 except KeyboardInterrupt: print('\nStopped')
 finally: server.server_close()
def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--workbook',type=Path); ap.add_argument('--output-root',type=Path,default=OUTPUT); ap.add_argument('--port',type=int,default=8000); ap.add_argument('--build-only',action='store_true'); a=ap.parse_args(); workbook=(a.workbook.resolve() if a.workbook else latest(a.output_root.resolve())); print('Workbook:',workbook)
 book=load_workbook(workbook,read_only=True,data_only=True); required={'Vision 2030',*SHEETS}; missing=required-set(book.sheetnames)
 if missing: raise ValueError('Missing sheets: '+', '.join(sorted(missing)))
 template,_,_=model((TEMPLATES/'briefing.html').read_text(encoding='utf-8')); generated,summary,details,source_map,duplicates=build(template,book,workbook); validation=audit(generated,summary,details,source_map,duplicates); print(json.dumps(validation,indent=2))
 if not validation['auditPass']: raise RuntimeError('Validation failed. Site was not launched. Review generated audits.')
 for page in PAGES: inject(TEMPLATES/page,SITE/page,generated)
 # page equality check
 signatures=[]
 for page in PAGES:
  m,_,_=model((SITE/page).read_text(encoding='utf-8')); signatures.append(json.dumps(m,sort_keys=True,ensure_ascii=False))
 if len(set(signatures))!=1: raise RuntimeError('Page data models differ')
 print('All pages built from one certified data model.')
 if not a.build_only: serve(a.port)
if __name__=='__main__': main()
